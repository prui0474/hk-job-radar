# -*- coding: utf-8 -*-
"""报告导出：把库里的岗位与投递进展导成一份能发出去的文件。

三种格式，都不需要联网：

- **HTML**：单文件，样式与图表全部内联（手写 SVG，不引 CDN）。
  双击就能看，也能直接 Ctrl+P 打印成纸或 PDF——这是最不容易坏的一种。
- **PDF**：用项目本来就装好的 Chromium 把上面那份 HTML 打成 A4 PDF，
  不额外引 Python 库（PDF 库要么依赖系统字体、要么中文字体得自己塞）。
- **Excel**：`.xlsx`，四个工作表（概览 / 岗位明细 / 投递记录 / 状态流水）。

三条刻意的设计：

1. **空库也能导出**：报表是「这一刻的快照」，0 条就老实写 0 条。
   直接报错会让用户以为功能坏了，而他要的其实是「证明我这周没偷懒」。
2. **所有文本都转义**：职位名、公司名来自第三方网站，可能带 `<` 或 `&`。
   导出的 HTML 是让人双击打开的文件，不能因为一条恶意 JD 就变成脚本。
3. **不写死来源**：库里有几个站就导几个，与看板的下拉框口径一致。
"""

from __future__ import annotations

import html
import io
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from loguru import logger

from hk_job_radar.db import Database
from hk_job_radar.models import JobPosting, now_hk

#: 状态 → 中文。与 `workbench_ui.html` 里的 STATUS_LABEL 是同一套口径，
#: 测试会盯着「看板能点出的状态，这里都得有中文名」，免得导出的表里
#: 冒出 `interview` 这种原文。
STATUS_LABEL: dict[str, str] = {
    "new": "待处理",
    "reported": "待处理",       # 已推送但还没动手，对用户来说仍是待处理
    "applied": "已投",
    "interview": "约面",
    "closed": "已结束",
    "rejected": "不感兴趣",
    "stale": "已归档",
    "expired": "已过期",
    "interested": "感兴趣",
    "clicked": "已点击",
}

#: 「动过手」的状态。导出「投递记录」这页时只收这些。
TOUCHED_STATUSES: tuple[str, ...] = ("applied", "interview", "closed", "rejected")

STATUS_LABEL_DEFAULT = "其他"

#: 明细表默认导多少条。库大了以后全量导出会让 HTML 变几十 MB，
#: 而「最近 500 条按分数排」已经覆盖了用户真正会翻的范围。
DEFAULT_LIMIT = 500


def _esc(value: Any) -> str:
    return html.escape(str(value if value is not None else ""), quote=True)


def _label(status: str) -> str:
    return STATUS_LABEL.get(status, STATUS_LABEL_DEFAULT)


def _day(value: Optional[datetime]) -> str:
    return value.strftime("%Y-%m-%d") if value else ""


def _num(value: Optional[int]) -> str:
    return "" if value is None else str(value)


def _pct(part: int, whole: int) -> str:
    """百分比；分母为 0 时返回破折号。

    不返回 0%：那是「投了很多但一个回复都没有」和「一个都没投」两种完全不同
    的处境，混成一个数字会误导人。
    """
    if whole <= 0:
        return "—"
    return f"{part / whole * 100:.0f}%"


# ---------------------------------------------------------------------------
# 取数
# ---------------------------------------------------------------------------
def job_rows(db: Database, *, limit: int = DEFAULT_LIMIT,
             statuses: tuple[str, ...] | None = None) -> list[dict[str, Any]]:
    """把岗位拍平成「一行一条」的字典，HTML 与 Excel 共用同一份。"""
    jobs = db.board_jobs(statuses=statuses, limit=limit)
    rows: list[dict[str, Any]] = []
    for job in jobs:
        rows.append({
            "uid": job.job_uid,
            "score": job.llm_score,
            "title": job.title,
            "company": job.company,
            "location": job.location,
            "salary": job.salary_text,
            "source": job.source,
            "status": _label(job.status),
            "status_raw": job.status,
            "posted": _day(job.posted_at or job.first_seen_at),
            "first_seen": _day(job.first_seen_at),
            "eligibility": job.eligibility_note,
            "reason": job.llm_reason,
            "keyword": job.keyword_matched,
            "url": job.url,
            "direct": "是" if job.direct_apply else "",
        })
    return rows


def status_events(db: Database, *, limit: int = 1000) -> list[dict[str, Any]]:
    """状态流水（最近 limit 条，新的在前），带上岗位标题方便直接看。"""
    if not hasattr(db, "status_history_rows"):
        return []
    try:
        events = db.status_history_rows(limit=limit)
    except Exception as exc:                       # noqa: BLE001 - 导出不该因此失败
        logger.warning(f"读状态流水失败：{exc}")
        return []
    cache: dict[str, JobPosting] = {}
    for row in events:
        uid = str(row.get("job_uid") or "")
        if uid and uid not in cache:
            job = db.get_job(uid)
            if job is not None:
                cache[uid] = job
        job = cache.get(uid)
        row["title"] = job.title if job else ""
        row["company"] = job.company if job else ""
        row["from"] = _label(str(row.get("from_status") or ""))
        row["to"] = _label(str(row.get("to_status") or ""))
    return events


def collect(db: Database, *, days: int = 30,
            limit: int = DEFAULT_LIMIT) -> dict[str, Any]:
    """导出需要的全部数据，一次取齐。

    分成两块：`insights` 来自看板进度页那套统计（口径直接复用，避免
    「报表说投了 3 个、看板说 2 个」这种事），`jobs` 是明细。
    """
    insights = db.board_insights(days=days)
    return {
        "generated_at": now_hk(),
        "days": days,
        "insights": insights,
        "applied": job_rows(db, limit=limit, statuses=TOUCHED_STATUSES),
        "jobs": job_rows(db, limit=limit),
        "events": status_events(db),
    }


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------
_CSS = """
:root{
  --bg:#f6f7f9; --card:#fff; --text:#1c2126; --muted:#6b7280; --line:#e3e6ea;
  --accent:#2f6feb; --ok:#0f9d58; --warn:#b45309; --danger:#d93025;
}
*{box-sizing:border-box}
body{margin:0; padding:32px 28px 48px; background:var(--bg); color:var(--text);
  font:14px/1.65 -apple-system,"Segoe UI","Microsoft YaHei","PingFang SC",sans-serif}
.wrap{max-width:1060px; margin:0 auto}
header{margin-bottom:22px}
h1{font-size:22px; margin:0 0 6px; letter-spacing:.2px}
.sub{color:var(--muted); font-size:13px}
.kpis{display:grid; grid-template-columns:repeat(4,1fr); gap:12px; margin:18px 0 22px}
.kpi{background:var(--card); border:1px solid var(--line); border-radius:12px; padding:14px 16px}
.kpi .n{font-size:24px; font-weight:700; line-height:1.2}
.kpi .k{color:var(--muted); font-size:12px; margin-top:2px}
section{background:var(--card); border:1px solid var(--line); border-radius:12px;
  padding:18px 20px; margin-bottom:18px}
h2{font-size:15px; margin:0 0 14px}
h3{font-size:13px; margin:18px 0 8px; color:var(--muted); font-weight:600}
table{width:100%; border-collapse:collapse; font-size:13px}
th,td{text-align:left; padding:7px 10px; border-bottom:1px solid var(--line); vertical-align:top}
th{color:var(--muted); font-weight:600; font-size:12px; white-space:nowrap}
tbody tr:hover{background:#fafbfc}
td.num{white-space:nowrap; font-variant-numeric:tabular-nums}
a{color:var(--accent); text-decoration:none}
a:hover{text-decoration:underline}
.badge{display:inline-block; padding:1px 8px; border-radius:999px; font-size:12px;
  background:#eef1f5; color:#495057; white-space:nowrap}
.badge.ok{background:#e6f4ea; color:var(--ok)}
.badge.warn{background:#fef3e2; color:var(--warn)}
.gate{color:var(--warn); font-size:12px}
.reason{color:var(--muted); font-size:12px}
.empty{color:var(--muted); font-size:13px; padding:10px 0}
footer{color:var(--muted); font-size:12px; text-align:center; margin-top:26px; line-height:1.9}
svg{display:block; max-width:100%}
@media print{
  body{background:#fff; padding:0}
  section{break-inside:avoid; border-color:#ddd}
  a{color:inherit}
}
"""


def _svg_funnel(funnel: dict[str, int]) -> str:
    """水平漏斗。纯手写 SVG：不引图表库，断网/内网也能看。"""
    rows = [
        ("入库岗位", int(funnel.get("pool") or 0), "#9aa4b2"),
        ("已投", int(funnel.get("applied") or 0), "#2f6feb"),
        ("约面", int(funnel.get("interview") or 0), "#0f9d58"),
    ]
    bar_x, bar_w, row_h = 84, 430, 38
    peak = max(1, rows[0][1])
    body = []
    for i, (label, value, color) in enumerate(rows):
        y = i * row_h
        width = 0 if value <= 0 else max(6, bar_w * value / peak)
        body.append(
            f'<g transform="translate(0,{y})">'
            f'<text x="0" y="23" font-size="13" fill="#6b7280">{label}</text>'
            + (f'<rect x="{bar_x}" y="8" width="{width:.1f}" height="20" rx="6"'
               f' fill="{color}" opacity=".88"/>' if width else "")
            + f'<text x="{bar_x + bar_w + 10}" y="23" font-size="13"'
              f' font-weight="600" fill="#1c2126">{value}</text></g>')
    deliver = int(funnel.get("applied") or 0)
    top = int(funnel.get("pool") or 0)
    return (
        f'<svg viewBox="0 0 620 {len(rows) * row_h + 26}" width="620"'
        f' height="{len(rows) * row_h + 26}" role="img" aria-label="投递漏斗">'
        + "".join(body)
        + f'<text x="0" y="{len(rows) * row_h + 16}" font-size="12" fill="#6b7280">'
          f'投出率 {_pct(deliver, top)} · 约面率 {_pct(int(funnel.get("interview") or 0), deliver)}'
          f'</text></svg>'
    )


def _source_table(sources: list[dict[str, Any]]) -> str:
    if not sources:
        return '<p class="empty">还没有按来源统计的数据。</p>'
    rows = "".join(
        f"<tr><td>{_esc(s.get('source'))}</td>"
        f"<td class='num'>{int(s.get('pool') or 0)}</td>"
        f"<td class='num'>{int(s.get('applied') or 0)}</td>"
        f"<td class='num'>{int(s.get('interview') or 0)}</td>"
        f"<td class='num'>{'—' if s.get('reply_rate') is None else str(s.get('reply_rate')) + '%'}</td></tr>"
        for s in sources)
    return ("<table><thead><tr><th>来源</th><th>入库</th><th>已投</th>"
            "<th>约面</th><th>回复率</th></tr></thead><tbody>"
            f"{rows}</tbody></table>"
            '<p class="reason" style="margin-top:8px">回复率 = 约面 ÷ 已投。'
            '样本只有几个的时候这个数字没有参考价值，别拿它下结论。</p>')


def _job_table(rows: list[dict[str, Any]], *, empty_text: str,
               with_status: bool = True) -> str:
    if not rows:
        return f'<p class="empty">{_esc(empty_text)}</p>'
    head = ("<tr><th>分数</th><th>职位</th><th>公司</th><th>地点</th>"
            + ("<th>状态</th>" if with_status else "")
            + "<th>来源</th><th>发布</th></tr>")
    body = []
    for r in rows:
        gate = (f'<div class="gate">{_esc(r["eligibility"])}</div>'
                if r.get("eligibility") else "")
        title = (f'<a href="{_esc(r["url"])}" target="_blank"'
                 f' rel="noopener noreferrer">{_esc(r["title"])}</a>')
        body.append(
            f'<tr><td class="num">{_num(r["score"])}</td>'
            f"<td>{title}{gate}</td>"
            f'<td>{_esc(r["company"])}</td>'
            f'<td>{_esc(r["location"])}</td>'
            + (f'<td><span class="badge">{_esc(r["status"])}</span></td>'
               if with_status else "")
            + f'<td>{_esc(r["source"])}</td>'
              f'<td class="num">{_esc(r["posted"])}</td></tr>')
    return (f"<table><thead>{head}</thead><tbody>{''.join(body)}</tbody></table>")


def render_html(data: dict[str, Any], *, title: str = "香港岗位雷达 · 投递进展报告") -> str:
    """生成单文件 HTML。返回值是完整文档，直接写盘即可。"""
    ins = data.get("insights") or {}
    funnel = ins.get("funnel") or {}
    generated = data.get("generated_at") or now_hk()
    stamp = generated.strftime("%Y-%m-%d %H:%M")
    pool = int(funnel.get("pool") or 0)
    applied = int(funnel.get("applied") or 0)
    interview = int(funnel.get("interview") or 0)
    rejected = int(funnel.get("rejected") or 0)
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_esc(title)} · {stamp}</title>
<style>{_CSS}</style>
</head><body><div class="wrap">
<header>
  <h1>{_esc(title)}</h1>
  <div class="sub">生成时间 {stamp}（香港时间） · 数据来源：本机数据库
    · 近 {int(data.get('days') or 30)} 天统计口径</div>
</header>

<div class="kpis">
  <div class="kpi"><div class="n">{pool}</div><div class="k">入库岗位</div></div>
  <div class="kpi"><div class="n">{applied}</div><div class="k">已投</div></div>
  <div class="kpi"><div class="n">{interview}</div><div class="k">约面</div></div>
  <div class="kpi"><div class="n">{rejected}</div><div class="k">不感兴趣</div></div>
</div>

<section>
  <h2>投递漏斗</h2>
  {_svg_funnel(funnel)}
</section>

<section>
  <h2>按来源看转化</h2>
  {_source_table(ins.get('sources') or [])}
</section>

<section>
  <h2>投递记录（{len(data.get('applied') or [])} 条）</h2>
  {_job_table(data.get('applied') or [], empty_text='还没有点过「已投」。')}
</section>

<section>
  <h2>岗位明细（{len(data.get('jobs') or [])} 条，按分数从高到低）</h2>
  {_job_table(data.get('jobs') or [], empty_text='库里还没有岗位。先跑一轮抓取。')}
</section>

<footer>
  这份报告由香港岗位雷达在本机生成，不含任何联网上报。<br>
  分数与理由是模型按你的画像给的参考意见，最终投不投永远由你自己判断。
</footer>
</div></body></html>
"""


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def build_xlsx(data: dict[str, Any]) -> bytes:
    """生成 .xlsx 的字节内容（四个工作表）。

    用 `BytesIO` 而不是先写临时文件：看板里点「导出 Excel」时直接把内存里的
    字节流回给浏览器，不留下需要清理的中间文件。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F6FEB")
    wrap = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()

    # --- 概览 ---
    ws = wb.active
    ws.title = "概览"
    ins = data.get("insights") or {}
    funnel = ins.get("funnel") or {}
    stamp = (data.get("generated_at") or now_hk()).strftime("%Y-%m-%d %H:%M")
    ws["A1"] = "香港岗位雷达 · 投递进展报告"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"生成时间：{stamp}（香港时间）"
    ws["A3"] = f"统计口径：近 {int(data.get('days') or 30)} 天"
    ws.append([])
    ws.append(["项目", "数量"])
    for cell in ws[5]:
        cell.font = head_font
        cell.fill = head_fill
    for key, label in (("pool", "入库岗位"), ("applied", "已投"),
                       ("interview", "约面"), ("rejected", "不感兴趣")):
        ws.append([label, int(funnel.get(key) or 0)])
    ws.append([])
    ws.append(["来源", "入库", "已投", "约面", "回复率"])
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = head_font
        cell.fill = head_fill
    for s in ins.get("sources") or []:
        rate = s.get("reply_rate")
        ws.append([s.get("source"), int(s.get("pool") or 0), int(s.get("applied") or 0),
                   int(s.get("interview") or 0),
                   "—" if rate is None else f"{rate}%"])
    ws.append([])
    ws.append(["说明：回复率 = 约面 ÷ 已投；样本很少时没有参考价值。"])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12

    # --- 岗位明细 / 投递记录 共用一个写表函数 ---
    def sheet_from_rows(name: str, rows: list[dict[str, Any]], *,
                        with_status: bool = True) -> None:
        ws = wb.create_sheet(name)
        header = ["分数", "职位", "公司", "地点", "薪资", "来源"]
        if with_status:
            header.append("状态")
        header += ["发布日", "入库日", "门槛提示", "AI 理由", "链接"]
        ws.append(header)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
        for r in rows:
            line = [r.get("score"), r.get("title"), r.get("company"),
                    r.get("location"), r.get("salary"), r.get("source")]
            if with_status:
                line.append(r.get("status"))
            line += [r.get("posted"), r.get("first_seen"),
                     r.get("eligibility"), r.get("reason"), r.get("url")]
            ws.append(line)
            url_cell = ws.cell(row=ws.max_row, column=len(header))
            if r.get("url"):
                url_cell.hyperlink = str(r["url"])
                url_cell.style = "Hyperlink"
                url_cell.value = "打开"
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
        widths = {"A": 7, "B": 42, "C": 26, "D": 16, "E": 18, "F": 10,
                  "G": 10, "H": 12, "I": 12, "J": 30, "K": 46, "L": 8}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap

    sheet_from_rows("岗位明细", data.get("jobs") or [])
    sheet_from_rows("投递记录", data.get("applied") or [])

    # --- 状态流水 ---
    ws = wb.create_sheet("状态流水")
    ws.append(["时间", "职位", "公司", "从", "到", "备注"])
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
    for e in data.get("events") or []:
        ws.append([str(e.get("created_at") or "")[:19], e.get("title") or "",
                   e.get("company") or "", e.get("from") or "",
                   e.get("to") or "", e.get("note") or ""])
    ws.freeze_panes = "A2"
    if data.get("events"):
        ws.auto_filter.ref = f"A1:F{ws.max_row}"
    for col, width in {"A": 20, "B": 42, "C": 26, "D": 12, "E": 12, "F": 24}.items():
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
class ExportError(RuntimeError):
    """导出失败，消息直接给用户看（中文、说清下一步怎么做）。"""


def build_pdf(html_path: Path, out_path: Path) -> Path:
    """用项目自带的 Chromium 把已经落盘的 HTML 打成 A4 PDF。

    为什么不用 Python 的 PDF 库：中文字体得自己塞进包里，体积大还容易缺字；
    而 Chromium 我们已经为了抓取装好了，`page.pdf()` 一次搞定，字形与
    打印效果就是浏览器「另存为 PDF」的那个结果。

    入参是**文件路径**而不是 HTML 字符串，因为 `file://` 打开时浏览器的
    字体回退、打印样式与用户自己双击看到的最接近。
    """
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise ExportError("没装 playwright，打不了 PDF。可以改用 --format html，"
                          "然后在浏览器里按 Ctrl+P 打印成 PDF。") from exc

    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(args=["--no-proxy-server"])
            try:
                page = browser.new_page()
                page.goto(Path(html_path).as_uri(), wait_until="load")
                page.pdf(path=str(out_path), format="A4", print_background=True,
                         margin={"top": "14mm", "bottom": "14mm",
                                 "left": "12mm", "right": "12mm"})
            finally:
                browser.close()
    except ExportError:
        raise
    except Exception as exc:                     # noqa: BLE001
        raise ExportError(
            f"PDF 生成失败：{exc}\n"
            "如果是第一次用浏览器，先运行 `playwright install chromium`；"
            "或者直接用 --format html，浏览器里 Ctrl+P 也能存成 PDF。") from exc
    return out_path


# ---------------------------------------------------------------------------
# 对外入口
# ---------------------------------------------------------------------------
def default_stem(when: Optional[datetime] = None) -> str:
    return "hk-job-radar-" + (when or now_hk()).strftime("%Y-%m-%d")


def export(db: Database, out_dir: Path, *,
           formats: tuple[str, ...] = ("html", "xlsx", "pdf"),
           days: int = 30, limit: int = DEFAULT_LIMIT,
           ) -> tuple[list[Path], list[str]]:
    """导出指定格式，返回 `(写出来的文件, 失败说明)`。

    单个格式失败**不牵连**其他格式：PDF 可能因浏览器环境打不出来，但 HTML
    和 Excel 已经写好了。用户不该因为一个格式失败就整次白跑——尤其是他要
    的往往是「能发出去的那一份」，HTML 就够。
    """
    data = collect(db, days=days, limit=limit)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = default_stem()
    written: list[Path] = []
    problems: list[str] = []

    if "html" in formats:
        path = out_dir / f"{stem}.html"
        path.write_text(render_html(data), encoding="utf-8")
        written.append(path)

    if "xlsx" in formats:
        if not xlsx_available():
            problems.append("Excel 没导出来：这台机器没装 openpyxl。装上就好——\n"
                            "  .venv\\Scripts\\pip install openpyxl")
        else:
            path = out_dir / f"{stem}.xlsx"
            path.write_bytes(build_xlsx(data))
            written.append(path)

    if "pdf" in formats:
        # PDF 由 HTML 打出来，所以就算只要 PDF 也得先有 HTML
        src = out_dir / f"{stem}.html"
        if not src.exists():
            src.write_text(render_html(data), encoding="utf-8")
        try:
            written.append(build_pdf(src, out_dir / f"{stem}.pdf"))
        except ExportError as exc:
            problems.append(str(exc))

    if not written and problems:
        raise ExportError("\n\n".join(problems))
    return written, problems
